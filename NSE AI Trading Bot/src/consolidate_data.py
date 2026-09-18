"""
consolidate_data.py
---------------------
Converts the per-ticker CSV files in data/ into a SINGLE multi-sheet XLSX
workbook (data/nifty_data.xlsx) -- one sheet per ticker, plus a combined
long-format sheet and a data-source/provenance sheet. CSV itself has no
concept of "sheets" (it is a flat, single-table format); XLSX is the
correct file format for that request. This script is the canonical way
to (re)produce the workbook from the per-ticker CSVs, and real_data_loader.py
now calls it automatically after ingestion.
"""

import os
import hashlib
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from universe import TICKERS

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE, "data")
WORKBOOK_PATH = os.path.join(DATA_DIR, "nifty_data.xlsx")

HEADER_FILL = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)


def format_sheet(ws, df):
    ws.freeze_panes = "A2"
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        max_len = max(len(str(col_name)), df[col_name].astype(str).map(len).max() if len(df) else 10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 22)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = BODY_FONT


def main():
    sheets = {}
    for t in TICKERS:
        sheets[t] = pd.read_csv(os.path.join(DATA_DIR, f"{t}.csv"), parse_dates=["date"])
    sheets["NIFTY50_INDEX"] = pd.read_csv(os.path.join(DATA_DIR, "NIFTY50_INDEX.csv"), parse_dates=["date"])
    sheets["Combined_LongFormat"] = pd.read_csv(os.path.join(DATA_DIR, "nifty_universe_combined.csv"), parse_dates=["date"])

    # Data-source / provenance summary sheet, built directly from the ticker
    # data itself (not parsed from DATA_SOURCE.txt, which is written AFTER
    # this workbook so it can include the workbook's own checksum -- avoids
    # a chicken-and-egg ordering problem).
    from universe import SECTOR_MAP
    source_rows = []
    for t in TICKERS:
        df = sheets[t]
        source_rows.append({
            "Ticker": t,
            "Sector": SECTOR_MAP.get(t, ""),
            "Start Date": df["date"].min().date(),
            "End Date": df["date"].max().date(),
            "Bars": len(df),
            "Close Min": round(df["close"].min(), 2),
            "Close Max": round(df["close"].max(), 2),
        })
    source_rows.append({
        "Ticker": "NIFTY50_INDEX", "Sector": "Broad Market Index",
        "Start Date": sheets["NIFTY50_INDEX"]["date"].min().date(),
        "End Date": sheets["NIFTY50_INDEX"]["date"].max().date(),
        "Bars": len(sheets["NIFTY50_INDEX"]),
        "Close Min": round(sheets["NIFTY50_INDEX"]["close"].min(), 2),
        "Close Max": round(sheets["NIFTY50_INDEX"]["close"].max(), 2),
    })
    sheets["DataSource_Info"] = pd.DataFrame(source_rows)

    with pd.ExcelWriter(WORKBOOK_PATH, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)  # Excel sheet-name limit: 31 chars

    from openpyxl import load_workbook
    wb = load_workbook(WORKBOOK_PATH)
    for name, df in sheets.items():
        ws = wb[name[:31]]
        format_sheet(ws, df)
    # Put the most-used sheets first for convenience
    order = [t[:31] for t in TICKERS] + ["NIFTY50_INDEX", "Combined_LongFormat", "DataSource_Info"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(WORKBOOK_PATH)

    checksum = hashlib.sha256(open(WORKBOOK_PATH, "rb").read()).hexdigest()
    print(f"Wrote consolidated workbook -> {WORKBOOK_PATH}")
    print(f"Sheets: {list(sheets.keys())}")
    print(f"SHA-256: {checksum}")
    return checksum


if __name__ == "__main__":
    main()
