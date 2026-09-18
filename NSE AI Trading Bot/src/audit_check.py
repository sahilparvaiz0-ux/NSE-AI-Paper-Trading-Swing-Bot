"""Final submission integrity and cross-outcome alignment audit."""
import os, sys, json, ast, zipfile
import pandas as pd
from openpyxl import load_workbook
from docx import Document
sys.path.insert(0, os.path.dirname(__file__))
from universe import TICKERS
BASE=os.path.dirname(os.path.dirname(os.path.abspath(__file__))); DATA=os.path.join(BASE,'data'); LOGS=os.path.join(BASE,'logs'); REPORT=os.path.join(BASE,'TradingBot_Report_Final.docx')
checks=[]
def ok(name,cond,detail=''):
 checks.append((name,bool(cond),detail)); print(('PASS' if cond else 'FAIL')+' | '+name+' | '+detail)
# workbook
path=os.path.join(DATA,'nifty_data.xlsx'); wb=load_workbook(path,read_only=True,data_only=True)
required=set(TICKERS+['All_Stocks','Data_Summary','Source_Notes','Validation','NIFTY50_INDEX'])
ok('Workbook sheets',required.issubset(set(wb.sheetnames)))
for t in TICKERS:
 df=pd.read_excel(path,sheet_name=t); dates=pd.to_datetime(df['Date'])
 ok(f'{t} rows/date range',len(df)==1241 and dates.min()==pd.Timestamp('2020-10-01') and dates.max()==pd.Timestamp('2025-09-30'),f'{len(df)} rows {dates.min().date()}->{dates.max().date()}')
# baseline outputs
perf=json.load(open(os.path.join(LOGS,'performance_summary.json'))); tr=pd.read_csv(os.path.join(LOGS,'trade_log.csv')); eq=pd.read_csv(os.path.join(LOGS,'equity_curve.csv'),parse_dates=['date'])
ok('Baseline ending equity',abs(perf['ending_equity']-1534921.9286)<1 and abs(eq.mark_to_market_equity.iloc[-1]-perf['ending_equity'])<1)
ok('Baseline total return',abs(perf['total_return_pct']-53.49219286)<.01)
ok('Baseline CAGR',abs(perf['cagr_pct']-8.95389222)<.01)
ok('Baseline Sharpe',abs(perf['sharpe_ratio']-.3530499)<.01)
ok('Baseline max drawdown',abs(perf['max_drawdown_pct']+7.740682)<.01)
ok('Trade count',len(tr)==98 and perf['total_trades']==98)
ok('Cash ledger non-negative',eq.cash.min()>=-0.01,str(eq.cash.min()))
ok('Trade P&L + interest reconciliation',abs(tr.pnl.sum()+perf['cash_interest_earned_inr']-(perf['ending_equity']-perf['starting_equity']))<1,str(tr.pnl.sum()+perf['cash_interest_earned_inr']))
# risk analysis
adv=json.load(open(os.path.join(LOGS,'integrated_analysis_summary.json'))); r=adv['risk_analytics']; ok('Risk VaR/CVaR alignment',abs(r['historical_var_95_daily_pct']+0.6225655)<.01 and abs(r['historical_cvar_95_daily_pct']+1.0413074)<.01)
ok('Return attribution alignment',abs(r['net_trading_pnl_inr']-tr.pnl.sum())<1 and abs(r['interest_on_idle_cash_inr']-perf['cash_interest_earned_inr'])<1)
# benchmark and advanced
bc=pd.read_csv(os.path.join(LOGS,'benchmark_comparison.csv')); ok('Vol-matched control present',any(bc.portfolio.str.contains('vol-matched')))
reg=json.load(open(os.path.join(LOGS,'market_model_regression.json'))); ok('Market regression alignment',abs(reg['beta_vs_equal_weight_basket']-.258606)<.01 and abs(reg['alpha_annualised_pct']+.226684)<.02)
null=json.load(open(os.path.join(LOGS,'monte_carlo_null.json'))); ok('Null model alignment',null['n_simulations']==2000 and abs(null['actual_percentile_within_null']-48.3)<.2 and abs(null['one_sided_p_value']-.517)<.01)
block=json.load(open(os.path.join(LOGS,'block_bootstrap.json'))); ok('Block bootstrap alignment',block['n_bootstrap_samples']==2000 and block['block_length']==35 and block['sharpe_ci_low']<0<block['sharpe_ci_high'])
cost=json.load(open(os.path.join(LOGS,'cost_decomposition.json'))); ok('Cost reconciliation',abs(cost['gross_pnl_inr']-cost['net_trading_pnl_inr']-cost['total_transaction_cost_inr'])<1 and abs(cost['total_transaction_cost_inr']-76340.1836)<1)
# sensitivity/OOS/AI
sens=pd.read_csv(os.path.join(LOGS,'sensitivity_analysis.csv')); ok('Sensitivity grid 20',len(sens)==20)
best=sens.loc[sens.sharpe_ratio.idxmax()]; ok('Best in-sample parameters',int(best.ema_fast)==10 and int(best.ema_slow)==30 and int(best.stop_atr_mult)==3 and int(best.target_atr_mult)==3 and abs(best.sharpe_ratio-.92)<.02)
oos=pd.read_csv(os.path.join(LOGS,'walk_forward_oos_result.csv'),header=None,index_col=0)[1]; ok('OOS result',abs(float(oos['oos_total_return_pct'])+8.4819228)<.01 and int(oos['oos_total_trades'])==41)
ai=json.load(open(os.path.join(LOGS,'ai_oos_result.json'))); ok('AI OOS result',abs(ai['auc']-.5157429)<.01 and abs(ai['ai_filtered_oos_total_return_pct']+4.6756713)<.01 and ai['ai_filtered_oos_total_trades']==34)
# report
D=Document(REPORT); txt='\n'.join(p.text for p in D.paragraphs); tables=['\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows) for t in D.tables]; alltxt=txt+'\n'+'\n'.join(tables)
for needle in ['53.49%','8.95%','-7.74%','-8.48%','-4.68%','0.516','-0.62%','-1.04%','48.3rd percentile','0.517','2,30,249','3,04,673','74.29%','0.26','-0.23%','6.94%']:
 ok('Report contains '+needle,needle in alltxt)
ok('No contradictory starting-equity note','approximately ₹10,00,258' not in alltxt and 'approximately Rs 10,00,258' not in alltxt)
ok('Advanced module documented','advanced_analysis.py' in alltxt)
ok('Business implications included','10.2 Business Implications' in txt)
# syntax and submission files
for fn in os.listdir(os.path.join(BASE,'src')):
 if fn.endswith('.py'):
  try: ast.parse(open(os.path.join(BASE,'src',fn),encoding='utf-8').read()); good=True; detail=''
  except Exception as e: good=False; detail=str(e)
  ok('Syntax '+fn,good,detail)
for f in ['README.md','requirements.txt','data/nifty_data.xlsx','logs/performance_summary.json','logs/trade_log.csv','logs/equity_curve.csv','logs/ai_oos_result.json','logs/monte_carlo_null.json','logs/block_bootstrap.json','TradingBot_Report_Final.docx']:
 ok('File '+f,os.path.exists(os.path.join(BASE,f)))
# no caches/temp backup
bad=[]
for root,dirs,files in os.walk(BASE):
 dirs[:] = [d for d in dirs if d not in ['__pycache__']]
 for f in files:
  if f.endswith(('.pyc','.tmp')) or f.startswith('~$'): bad.append(os.path.join(root,f))
ok('No temp/compiled files',not bad,str(bad))
passed=sum(c for _,c,_ in checks); total=len(checks)
with open(os.path.join(LOGS,'audit_check.txt'),'w') as f:
 f.write(f'FINAL AUDIT: {passed}/{total} checks passed\n')
 for n,c,d in checks:f.write(('PASS' if c else 'FAIL')+f' | {n} | {d}\n')
print(f'FINAL AUDIT: {passed}/{total} checks passed')
if passed!=total: sys.exit(1)
